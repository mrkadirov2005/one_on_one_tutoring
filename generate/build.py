# -*- coding: utf-8 -*-
"""
Generates lesson.html and slides.pptx for every lesson in lessons_data.py,
plus a root index.html course overview.

Usage:
    python3 generate/build.py
Run from the project root (/Users/furb-x/Desktop/new_lesson).
"""

import os
import html
import sys

sys.path.insert(0, os.path.dirname(__file__))
from lessons_data import LESSONS

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BY_ID = {l["id"]: l for l in LESSONS}

GREEN = RGBColor(0x2C, 0x5F, 0x2D)
GREEN_LIGHT = RGBColor(0xEE, 0xF6, 0xEE)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
ORANGE = RGBColor(0xB4, 0x53, 0x09)
DARK = RGBColor(0x1E, 0x25, 0x30)
LIGHT_TXT = RGBColor(0xE6, 0xE6, 0xE6)
INK = RGBColor(0x22, 0x22, 0x22)


# ------------------------------------------------------------------ HTML ---

def e(text):
    return html.escape(text, quote=False)


def when_label(lesson):
    if lesson["week"] == 0:
        return f"Before Week 1 · {lesson['day']}"
    return f"Week {lesson['week']} · {lesson['day']}"


def order_label(lesson):
    if lesson["order"] == 0:
        return "Introduction"
    return f"Lesson {lesson['order']} of 12"


def render_lesson_html(lesson):
    checks = BY_ID.get(lesson["checks_hw_from"]) if lesson["checks_hw_from"] else None

    parts = []
    parts.append("<!DOCTYPE html>")
    parts.append('<html lang="en">')
    parts.append("<head>")
    parts.append('<meta charset="UTF-8">')
    parts.append('<meta name="viewport" content="width=device-width, initial-scale=1">')
    parts.append(f"<title>{e(lesson['topic'])}</title>")
    parts.append('<link rel="stylesheet" href="../../style.css">')
    parts.append("</head>")
    parts.append("<body>")

    parts.append('<div class="top-nav"><a href="../../index.html">&larr; Course overview</a></div>')

    parts.append(
        f"<h1>{e(lesson['topic'])}"
        f"<span class=\"day-badge\">{e(when_label(lesson))}</span></h1>"
    )

    if checks:
        parts.append(
            f'<p><span class="tag-review">Homework check</span> Today we review the homework '
            f'assigned in <em>{e(checks["topic"])}</em> ({checks["day"]}, Week {checks["week"]}).</p>'
        )

    # Description
    parts.append("<section>")
    parts.append("<h2>Description</h2>")
    parts.append(f"<p>{e(lesson['description'])}</p>")
    parts.append("</section>")

    # Key concepts
    parts.append("<section>")
    parts.append("<h2>Key Concepts</h2>")
    parts.append('<ul class="concepts">')
    for c in lesson["key_concepts"]:
        parts.append(f"<li>{e(c)}</li>")
    parts.append("</ul>")
    parts.append("</section>")

    for extra in lesson.get("extra_sections", []):
        parts.append("<section>")
        parts.append(f"<h2>{e(extra['title'])}</h2>")
        parts.append('<ul class="concepts">')
        for item in extra["items"]:
            parts.append(f"<li>{e(item)}</li>")
        parts.append("</ul>")
        parts.append("</section>")

    # Code snippets
    parts.append("<section>")
    parts.append("<h2>Code Samples</h2>")
    for s in lesson["snippets"]:
        parts.append('<div class="snippet">')
        parts.append(f'<p class="def"><strong>What it does:</strong> {e(s["def"])}</p>')
        parts.append(f'<pre><code>{e(s["code"])}</code></pre>')
        parts.append("</div>")
    parts.append("</section>")

    # Example
    parts.append("<section>")
    parts.append("<h2>Example</h2>")
    parts.append('<div class="example">')
    parts.append(f'<p>{e(lesson["example"]["desc"])}</p>')
    parts.append(f'<pre><code>{e(lesson["example"]["code"])}</code></pre>')
    parts.append("</div>")
    parts.append("</section>")

    # Exercises
    hw_note = ' <span class="tag-review">Homework — due next Wednesday</span>' if lesson["homework"] else ""
    parts.append('<div class="exercises">')
    parts.append(f"<h2>Exercises{hw_note}</h2>")
    parts.append('<ul class="exercises">')
    for ex in lesson["exercises"]:
        parts.append(f"<li>{e(ex)}</li>")
    parts.append("</ul>")
    parts.append("</div>")

    parts.append(f"<footer>{order_label(lesson)} &middot; {e(lesson['topic'])}</footer>")
    parts.append("</body>")
    parts.append("</html>")
    return "\n".join(parts)


# ------------------------------------------------------------------ PPTX ---

BLANK = 6  # blank layout index in the default template


def add_title_bar(slide, prs, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), prs.slide_width - Inches(1.2), Inches(1.3))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = GREEN
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(18)
        r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return box


def add_bullets(slide, prs, items, top=Inches(1.6), font_size=24, color=INK, bold_first=False):
    box = slide.shapes.add_textbox(Inches(0.8), top, prs.slide_width - Inches(1.6), prs.slide_height - top - Inches(0.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        p.space_after = Pt(10)
    return box


def add_paragraph_text(slide, prs, text, top=Inches(1.7), font_size=22):
    box = slide.shapes.add_textbox(Inches(0.8), top, prs.slide_width - Inches(1.6), prs.slide_height - top - Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = INK
    return box


def add_code_block(slide, prs, code, top=Inches(1.6), height=None, font_size=16):
    left = Inches(0.7)
    width = prs.slide_width - Inches(1.4)
    if height is None:
        height = prs.slide_height - top - Inches(0.5)
    shape = slide.shapes.add_shape(1, left, top, width, height)  # 1 = rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.15)
    lines = code.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line if line.strip() else " "
        run.font.name = "Menlo"
        run.font.size = Pt(font_size)
        run.font.color.rgb = LIGHT_TXT
    return shape


def add_caption(slide, prs, text, top, color=BLUE):
    box = slide.shapes.add_textbox(Inches(0.7), top, prs.slide_width - Inches(1.4), Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(17)
    run.font.italic = True
    run.font.color.rgb = color
    return box


def build_lesson_pptx(lesson, out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[BLANK])

    # 1. Title slide
    s = new_slide()
    slide_center_title(s, prs, lesson["topic"], f"{when_label(lesson)} · {order_label(lesson)}")

    # 2. Description
    s = new_slide()
    add_title_bar(s, prs, "Description")
    add_paragraph_text(s, prs, lesson["description"])

    # 3. Key concepts
    s = new_slide()
    add_title_bar(s, prs, "Key Concepts")
    add_bullets(s, prs, lesson["key_concepts"])

    for extra in lesson.get("extra_sections", []):
        s = new_slide()
        add_title_bar(s, prs, extra["title"])
        add_bullets(s, prs, extra["items"], font_size=21)

    # 4..N Code snippets (one slide each)
    for i, snip in enumerate(lesson["snippets"], start=1):
        s = new_slide()
        add_title_bar(s, prs, f"Code Sample {i} of {len(lesson['snippets'])}")
        add_caption(s, prs, snip["def"], top=Inches(1.5))
        add_code_block(s, prs, snip["code"], top=Inches(2.15))

    # Example
    s = new_slide()
    add_title_bar(s, prs, "Example")
    add_caption(s, prs, lesson["example"]["desc"], top=Inches(1.5))
    add_code_block(s, prs, lesson["example"]["code"], top=Inches(2.15))

    # Exercises
    s = new_slide()
    title = "Exercises"
    if lesson["homework"]:
        title += "  (Homework — due next Wednesday)"
    add_title_bar(s, prs, title)
    add_bullets(s, prs, lesson["exercises"], font_size=22, color=ORANGE if lesson["homework"] else INK)

    prs.save(out_path)


def slide_center_title(slide, prs, title, subtitle):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), prs.slide_width - Inches(1.6), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = GREEN

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = subtitle
    r2.font.size = Pt(20)
    r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return box


# ------------------------------------------------------------- ROOT INDEX --

def render_index_html():
    weeks = {}
    for l in LESSONS:
        weeks.setdefault(l["week"], []).append(l)

    week_titles = {
        0: "Introduction & Setup",
        1: "HTML Foundations",
        2: "HTML Wrap-up + CSS Intro",
        3: "CSS Deep Dive",
        4: "JavaScript Basics",
    }

    parts = []
    parts.append("<!DOCTYPE html>")
    parts.append('<html lang="en">')
    parts.append("<head>")
    parts.append('<meta charset="UTF-8">')
    parts.append('<meta name="viewport" content="width=device-width, initial-scale=1">')
    parts.append("<title>HTML/CSS/JS Course — One Month Plan</title>")
    parts.append('<link rel="stylesheet" href="style.css">')
    parts.append("<style>")
    parts.append("table { width: 100%; border-collapse: collapse; margin: 10px 0 26px; }")
    parts.append("th, td { text-align: left; padding: 8px 10px; border-bottom: 1px solid #e2e2e2; }")
    parts.append("th { color: #2c5f2d; font-size: 0.85em; text-transform: uppercase; letter-spacing: 0.03em; }")
    parts.append("td a { color: #1d4ed8; text-decoration: none; }")
    parts.append("td a:hover { text-decoration: underline; }")
    parts.append(".hw-cell { color: #b45309; font-weight: 600; font-size: 0.85em; }")
    parts.append("</style>")
    parts.append("</head>")
    parts.append("<body>")
    parts.append("<h1>HTML, CSS &amp; JS in One Month</h1>")
    parts.append(
        "<p>An introduction class, then 12 lessons, 3 per week (Wednesday, Saturday, Sunday). Homework is assigned on "
        "Saturday and reviewed the following Wednesday. Click a lesson to open its page; "
        "each lesson folder also has a matching <code>slides.pptx</code> for class.</p>"
    )

    for wk in sorted(weeks):
        heading = e(week_titles[wk]) if wk == 0 else f"Week {wk} — {e(week_titles[wk])}"
        parts.append(f"<h2>{heading}</h2>")
        parts.append("<table>")
        parts.append("<tr><th>Day</th><th>Topic</th><th>Notes</th></tr>")
        for l in weeks[wk]:
            note = ""
            if l["homework"]:
                note = '<span class="hw-cell">Homework assigned</span>'
            elif l["checks_hw_from"]:
                note = '<span class="hw-cell">Homework check</span>'
            link = f"{l['folder']}/lesson.html"
            parts.append(
                f"<tr><td>{e(l['day'])}</td>"
                f"<td><a href=\"{link}\">{e(l['topic'])}</a></td>"
                f"<td>{note}</td></tr>"
            )
        parts.append("</table>")

    parts.append('<footer>Plan: <a href="plan/main_plan.html">plan/main_plan.html</a></footer>')
    parts.append("</body>")
    parts.append("</html>")
    return "\n".join(parts)


# --------------------------------------------------------------- RUNNER ----

def main():
    for lesson in LESSONS:
        folder = os.path.join(ROOT, lesson["folder"])
        os.makedirs(folder, exist_ok=True)

        html_path = os.path.join(folder, "lesson.html")
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(render_lesson_html(lesson))

        pptx_path = os.path.join(folder, "slides.pptx")
        build_lesson_pptx(lesson, pptx_path)

        print(f"built {lesson['folder']}/lesson.html + slides.pptx")

    index_path = os.path.join(ROOT, "index.html")
    with open(index_path, "w", encoding="utf-8") as f:
        f.write(render_index_html())
    print("built index.html")


if __name__ == "__main__":
    main()
